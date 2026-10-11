"""Result, flag and audit exports (FR-6.10.5, AC-12, spec section 14)."""

import csv
import io
from datetime import timedelta

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.services.sessions import start_session
from apps.audit.models import AuditEvent
from apps.core.domain.permissions import Perm, Role
from apps.core.ids import uuid7
from apps.core.rls import system_context
from apps.geography.models import PollingUnit
from apps.reports import services
from apps.reports.models import ExportJob

from ..anomalies.helpers import unit_with_coordinates

pytestmark = pytest.mark.django_db


def _post(client, body, key=None):
    return client.post(
        "/api/v1/exports", body, format="json", HTTP_IDEMPOTENCY_KEY=key or str(uuid7())
    )


def _job(job_id):
    with system_context():
        return ExportJob.objects.get(id=job_id)


def _read(storage, job) -> bytes:
    return storage.read(job.object_key)


def _csv_rows(storage, job):
    text = _read(storage, job).decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


@pytest.fixture
def busy(field_):
    f = field_
    f.submit(f.pus_of(f.wards[0])[0], (100, 40))
    f.submit(f.pus_of(f.wards[0])[1], (20, 30))
    f.submit(f.pus_of(f.wards[2])[0], (7, 70))
    return f


def _results_body(f, **extra):
    return {
        "kind": "RESULTS",
        "format": "CSV",
        "election_id": str(f.election.id),
        "contest_id": str(f.contest.id),
        **extra,
    }


def test_a_results_export_is_built_in_the_background_and_ready_to_download(busy, storage):
    f = busy
    response = _post(f.admin_client(), _results_body(f))
    assert response.status_code == 202, response.content
    job = _job(response.json()["id"])
    assert job.status == "READY"  # the worker ran after commit
    assert job.row_count == 8 and job.size > 0 and len(job.sha256) == 64
    assert "PROVISIONAL" in job.filename and job.filename.endswith(".csv")

    rows = _csv_rows(storage, job)
    assert len(rows) == 8
    reported = [r for r in rows if r["status"] != "NO_RESULT"]
    assert len(reported) == 3
    assert sum(int(r["votes_AAA"]) for r in rows) == 127
    assert sum(int(r["votes_BBB"]) for r in rows) == 140
    assert {"lga_name", "ward_name", "pu_code", "valid", "leader", "margin"} <= set(rows[0])

    link = f.admin_client().post(f"/api/v1/exports/{job.id}/download-url")
    assert link.status_code == 200, link.content
    assert link.json()["sha256"] == job.sha256
    actions = set(AuditEvent.objects.filter(object_id=str(job.id)).values_list("action", flat=True))
    assert {"export.requested", "export.completed", "export.download_url_issued"} <= actions


def test_ward_and_lga_exports_roll_up(busy, storage):
    f = busy
    ward_job = _job(_post(f.admin_client(), _results_body(f, level="ward")).json()["id"])
    ward_rows = _csv_rows(storage, ward_job)
    assert len(ward_rows) == 4
    assert sum(int(r["reported_pus"]) for r in ward_rows) == 3
    lga_job = _job(_post(f.admin_client(), _results_body(f, level="lga")).json()["id"])
    lga_rows = _csv_rows(storage, lga_job)
    assert [(r["lga_name"], r["reported_pus"]) for r in lga_rows] == [
        ("Synthetic LGA 01", "2"),
        ("Synthetic LGA 02", "1"),
    ]


def test_the_verified_measure_counts_only_approved_results(reviewed_field, storage):
    f = reviewed_field
    approved = f.submit(f.pus[0], (60, 40))
    f.submit(f.pus[1], (10, 10))
    f.admin_client().post(
        f"/api/v1/results/{approved['version_id']}/review",
        {"decision": "APPROVED"},
        format="json",
        HTTP_IDEMPOTENCY_KEY=str(uuid7()),
    )
    job = _job(_post(f.admin_client(), _results_body(f, measure="verified")).json()["id"])
    rows = _csv_rows(storage, job)
    assert sum(int(r["votes_AAA"]) for r in rows) == 60


def test_xlsx_carries_the_notice(busy, storage):
    f = busy
    job = _job(_post(f.admin_client(), _results_body(f, format="XLSX")).json()["id"])
    book = load_workbook(io.BytesIO(_read(storage, job)))
    assert book.sheetnames == ["Results", "About"]
    about = dict(book["About"].iter_rows(min_row=2, values_only=True))
    assert "not an official declaration" in about["notice"].lower()
    assert about["rows"] == 8
    assert book["Results"].max_row == 9  # header plus eight units


def test_spreadsheet_formulas_in_data_are_neutralized(busy, storage):
    f = busy
    unit = f.pus_of(f.wards[0])[0]
    with system_context():
        PollingUnit.objects.filter(id=unit.id).update(name='=HYPERLINK("http://evil","click")')
    job = _job(_post(f.admin_client(), _results_body(f)).json()["id"])
    names = [r["pu_name"] for r in _csv_rows(storage, job)]
    assert '\'=HYPERLINK("http://evil","click")' in names
    assert not any(n.startswith("=") for n in names)


def test_exports_are_scoped_to_the_requesters_geography(busy, storage):
    f = busy
    member = f.make_member(
        Role.LGA_OFFICER,
        scope_ids=[f.lgas[0].id],
        organization=f.org,
        granted=[Perm.REPORTS_EXPORT],
        mfa_enabled=True,
    )
    job = _job(_post(f.client_as(member), _results_body(f)).json()["id"])
    rows = _csv_rows(storage, job)
    assert len(rows) == 4
    assert {r["lga_name"] for r in rows} == {"Synthetic LGA 01"}
    assert sum(int(r["votes_AAA"]) for r in rows) == 120


def test_who_may_request_an_export(busy):
    f = busy
    ward = f.client_as(f.officer(Role.WARD_OFFICER, f.wards[0]))
    assert _post(ward, _results_body(f)).status_code == 403
    agent = f.agents[f.pus[0].id].client
    assert _post(agent, _results_body(f)).status_code == 403
    plain_lga = f.client_as(f.officer(Role.LGA_OFFICER, f.lgas[0]))  # no export grant
    assert _post(plain_lga, _results_body(f)).status_code == 403


def test_a_stale_second_factor_is_refused(busy):
    f = busy
    stale = start_session(
        user=f.admin.user,
        membership=f.admin.membership,
        client="WEB",
        device=None,
        mfa_at=timezone.now() - timedelta(hours=2),
        ip=None,
    )
    client = f.client_as(f.admin)
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {stale.access_token}")
    response = _post(client, _results_body(f))
    assert (response.status_code, response.json()["code"]) == (403, "mfa_step_up_required")


def test_the_same_key_gives_the_same_job(busy):
    f = busy
    key = str(uuid7())
    first = _post(f.admin_client(), _results_body(f), key).json()
    second = _post(f.admin_client(), _results_body(f), key).json()
    assert first["id"] == second["id"]
    with system_context():
        assert ExportJob.objects.count() == 1


def test_people_cannot_run_many_exports_at_once(busy, settings):
    f = busy
    settings.VTRS_EXPORT_MAX_ACTIVE = 1
    with system_context():
        ExportJob.objects.create(
            organization=f.org,
            requested_by=f.admin.user,
            kind="RESULTS",
            format="CSV",
            actor_role="PARTY_ADMIN",
            scope_kind="ORG",
            status="RUNNING",
        )
    response = _post(f.admin_client(), _results_body(f))
    assert (response.status_code, response.json()["code"]) == (429, "too_many_exports")


@pytest.mark.parametrize(
    "mutate",
    [
        {"format": "ZIP"},  # not a results format
        {"level": "galaxy"},
        {"measure": "maybe"},
        {"contest_id": None},
    ],
)
def test_invalid_requests_are_refused(busy, mutate):
    f = busy
    body = {**_results_body(f), **mutate}
    body = {k: v for k, v in body.items() if v is not None}
    assert _post(f.admin_client(), body).status_code == 400


def test_a_contest_from_another_election_is_refused(busy):
    f = busy
    body = _results_body(f, contest_id=str(uuid7()))
    assert _post(f.admin_client(), body).status_code in {400, 404}


def test_other_people_cannot_see_or_download_my_export(busy):
    f = busy
    mine = _post(f.admin_client(), _results_body(f)).json()["id"]
    other = f.make_member(
        Role.LGA_OFFICER,
        scope_ids=[f.lgas[0].id],
        organization=f.org,
        granted=[Perm.REPORTS_EXPORT],
        mfa_enabled=True,
    )
    client = f.client_as(other)
    assert client.get(f"/api/v1/exports/{mine}").status_code == 404
    assert client.post(f"/api/v1/exports/{mine}/download-url").status_code == 404
    assert client.get("/api/v1/exports").json()["results"] == []
    assert len(f.admin_client().get("/api/v1/exports").json()["results"]) == 1


def test_downloads_need_a_fresh_second_factor_and_a_live_file(busy, storage):
    f = busy
    job = _job(_post(f.admin_client(), _results_body(f)).json()["id"])
    stale = start_session(
        user=f.admin.user,
        membership=f.admin.membership,
        client="WEB",
        device=None,
        mfa_at=timezone.now() - timedelta(hours=2),
        ip=None,
    )
    client = f.client_as(f.admin)
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {stale.access_token}")
    refused = client.post(f"/api/v1/exports/{job.id}/download-url")
    assert refused.json()["code"] == "mfa_step_up_required"

    with system_context():
        ExportJob.objects.filter(id=job.id).update(expires_at=timezone.now() - timedelta(minutes=1))
    expired = f.admin_client().post(f"/api/v1/exports/{job.id}/download-url")
    assert (expired.status_code, expired.json()["code"]) == (410, "export_expired")
    assert services.purge_expired() == 1
    assert storage.size_of(job.object_key) is None
    assert _job(job.id).status == "EXPIRED"
    again = f.admin_client().post(f"/api/v1/exports/{job.id}/download-url")
    assert again.status_code == 410
    assert AuditEvent.objects.filter(action="export.purged", object_id=str(job.id)).exists()


def test_a_job_that_is_not_finished_has_no_link(busy):
    f = busy
    with system_context():
        job = ExportJob.objects.create(
            organization=f.org,
            requested_by=f.admin.user,
            kind="RESULTS",
            format="CSV",
            actor_role="PARTY_ADMIN",
            scope_kind="ORG",
            status="PENDING",
        )
    response = f.admin_client().post(f"/api/v1/exports/{job.id}/download-url")
    assert (response.status_code, response.json()["code"]) == (409, "export_not_ready")


def test_a_failed_build_is_recorded_without_leaking_internals(busy, monkeypatch):
    f = busy

    def boom(*args, **kwargs):
        raise RuntimeError("secret internal path /srv/keys")

    monkeypatch.setattr("apps.reports.builders.build_results", boom)
    job = _job(_post(f.admin_client(), _results_body(f)).json()["id"])
    assert job.status == "FAILED"
    assert job.error == "The export could not be built."
    assert "secret" not in job.error
    assert AuditEvent.objects.filter(action="export.failed", object_id=str(job.id)).exists()


def test_a_flag_register_export(field_, storage):
    f = field_
    pu = unit_with_coordinates(f)
    f.submit(
        pu,
        gps={
            "lat": round(float(pu.latitude) + 0.5, 6),
            "lng": float(pu.longitude),
            "accuracy_m": 9,
        },
    )
    body = {"kind": "FLAGS", "format": "CSV", "election_id": str(f.election.id)}
    job = _job(_post(f.admin_client(), body).json()["id"])
    rows = _csv_rows(storage, job)
    assert [r["rule_code"] for r in rows] == ["GPS_OUT_OF_GEOFENCE"]
    assert rows[0]["severity"] == "HIGH"
    assert rows[0]["pu_code"] == pu.inec_code
    filtered = {**body, "severity": "CRITICAL"}
    assert _csv_rows(storage, _job(_post(f.admin_client(), filtered).json()["id"])) == []


def test_an_audit_export_is_windowed_and_admin_only(busy, storage):
    f = busy
    now = timezone.now()
    body = {
        "kind": "AUDIT",
        "format": "CSV",
        "since": (now - timedelta(hours=1)).isoformat(),
        "until": (now + timedelta(minutes=5)).isoformat(),
        "action": "result.",
    }
    job = _job(_post(f.admin_client(), body).json()["id"])
    rows = _csv_rows(storage, job)
    assert rows and all(r["action"].startswith("result.") for r in rows)
    too_long = {**body, "since": (now - timedelta(days=40)).isoformat()}
    assert _post(f.admin_client(), too_long).status_code == 400
    assert _post(f.admin_client(), {**body, "until": body["since"]}).status_code == 400
    ward = f.client_as(f.officer(Role.WARD_OFFICER, f.wards[0]))
    assert _post(ward, body).status_code == 403
